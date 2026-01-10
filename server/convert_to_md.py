"""
Script to convert .docx files to .md format for RAG chatbot system
"""
from docx import Document
import os
import re

def clean_text(text):
    """Clean up text for markdown format"""
    # Remove excessive whitespace
    text = re.sub(r'\n\s*\n\s*\n+', '\n\n', text)
    text = re.sub(r' +', ' ', text)
    return text.strip()

def convert_docx_to_md(docx_path, md_path):
    """Convert a .docx file to .md format"""
    try:
        # Load the document
        doc = Document(docx_path)
        
        md_content = []
        
        # Process each paragraph
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            
            if not text:
                md_content.append('')
                continue
            
            # Check if it's a heading based on style
            style_name = paragraph.style.name.lower()
            
            if 'heading 1' in style_name:
                md_content.append(f'# {text}')
            elif 'heading 2' in style_name:
                md_content.append(f'## {text}')
            elif 'heading 3' in style_name:
                md_content.append(f'### {text}')
            elif 'heading 4' in style_name:
                md_content.append(f'#### {text}')
            else:
                md_content.append(text)
        
        # Process tables if any
        for table in doc.tables:
            md_content.append('')  # Add spacing before table
            
            # Get headers
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            md_content.append('| ' + ' | '.join(headers) + ' |')
            md_content.append('| ' + ' | '.join(['---'] * len(headers)) + ' |')
            
            # Get rows
            for row in table.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                md_content.append('| ' + ' | '.join(cells) + ' |')
            
            md_content.append('')  # Add spacing after table
        
        # Join content and clean
        final_content = '\n'.join(md_content)
        final_content = clean_text(final_content)
        
        # Write to markdown file
        with open(md_path, 'w', encoding='utf-8') as f:
            f.write(final_content)
        
        print(f'✓ Converted: {os.path.basename(docx_path)} -> {os.path.basename(md_path)}')
        return True
        
    except Exception as e:
        print(f'✗ Error converting {os.path.basename(docx_path)}: {str(e)}')
        return False

def main():
    """Main function to convert all .docx files in chatbot-knowledge directory"""
    knowledge_dir = 'chatbot-knowledge'
    
    if not os.path.exists(knowledge_dir):
        print(f'Error: Directory {knowledge_dir} not found!')
        return
    
    # Get all .docx files
    docx_files = [f for f in os.listdir(knowledge_dir) if f.endswith('.docx')]
    
    if not docx_files:
        print('No .docx files found in chatbot-knowledge directory')
        return
    
    print(f'Found {len(docx_files)} .docx files to convert\n')
    
    success_count = 0
    
    # Convert each file
    for docx_file in docx_files:
        docx_path = os.path.join(knowledge_dir, docx_file)
        md_file = docx_file.replace('.docx', '.md')
        md_path = os.path.join(knowledge_dir, md_file)
        
        if convert_docx_to_md(docx_path, md_path):
            success_count += 1
    
    print(f'\n✓ Successfully converted {success_count}/{len(docx_files)} files')

if __name__ == '__main__':
    main()
